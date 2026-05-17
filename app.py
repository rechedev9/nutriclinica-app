from flask import Flask, request, jsonify, send_file, render_template, send_from_directory
import sqlite3, os, uuid, io, unicodedata, shutil
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def nombre_a_prefijo(nombre):
    n = unicodedata.normalize('NFKD', nombre).encode('ascii', 'ignore').decode('ascii')
    return '_'.join(n.lower().split())

def nombre_a_carpeta(nombre):
    n = unicodedata.normalize('NFKD', nombre).encode('ascii', 'ignore').decode('ascii')
    return ''.join(w.capitalize() for w in n.split())

app      = Flask(__name__)
app.config['TEMPLATES_AUTO_RELOAD'] = True
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH  = os.path.join(BASE_DIR, 'bbdd', 'pacientes.db')
FOTOS_DIR = os.path.join(BASE_DIR, 'fotos')
os.makedirs(FOTOS_DIR, exist_ok=True)
os.makedirs(os.path.join(BASE_DIR, 'bbdd'), exist_ok=True)


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    conn.execute('''CREATE TABLE IF NOT EXISTS pacientes (
        id               INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre           TEXT NOT NULL,
        fecha_nac        TEXT,
        sexo             TEXT,
        dni              TEXT,
        telefono         TEXT,
        email            TEXT,
        objetivo         TEXT,
        actividad_fisica TEXT,
        intolerancias    TEXT,
        patologias       TEXT,
        notas            TEXT,
        peso             REAL,
        talla            REAL,
        creado           TEXT DEFAULT (datetime('now','localtime'))
    )''')
    for col in ('peso', 'talla'):
        try:
            conn.execute(f'ALTER TABLE pacientes ADD COLUMN {col} REAL')
        except Exception:
            pass
    conn.execute('''CREATE TABLE IF NOT EXISTS revisiones (
        id              INTEGER PRIMARY KEY AUTOINCREMENT,
        paciente_id     INTEGER NOT NULL,
        fecha           TEXT DEFAULT (date('now','localtime')),
        peso            REAL,
        talla           REAL,
        envergadura     REAL,
        pl_bicipital    REAL,
        pl_tricipital   REAL,
        pl_subescapular REAL,
        pl_supraespinal REAL,
        pl_supracrestal REAL,
        pl_abdominal    REAL,
        pl_muslo_int    REAL,
        pl_pierna_med   REAL,
        per_antebrazo   REAL,
        per_brazo_rel   REAL,
        per_brazo_cont  REAL,
        per_hombros     REAL,
        per_pecho       REAL,
        per_cintura     REAL,
        per_caderas     REAL,
        per_muslo       REAL,
        per_muslo_cont  REAL,
        per_gemelo      REAL,
        notas           TEXT,
        FOREIGN KEY (paciente_id) REFERENCES pacientes(id)
    )''')
    conn.execute('''CREATE TABLE IF NOT EXISTS fotos (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        paciente_id INTEGER NOT NULL,
        filename    TEXT NOT NULL,
        nota        TEXT,
        fecha       TEXT DEFAULT (datetime('now','localtime')),
        FOREIGN KEY (paciente_id) REFERENCES pacientes(id)
    )''')
    conn.commit()
    conn.close()


# ─── RUTAS PRINCIPALES ────────────────────────────────────────────────────────

@app.route('/')
def index():
    return render_template('index.html')


# ─── PACIENTES ────────────────────────────────────────────────────────────────

@app.route('/api/pacientes', methods=['GET'])
def get_pacientes():
    conn = get_db()
    rows = conn.execute('SELECT * FROM pacientes ORDER BY nombre COLLATE NOCASE').fetchall()
    result = []
    for p in rows:
        cnt     = conn.execute('SELECT COUNT(*) FROM fotos WHERE paciente_id=?', (p['id'],)).fetchone()[0]
        rev_cnt = conn.execute('SELECT COUNT(*) FROM revisiones WHERE paciente_id=?', (p['id'],)).fetchone()[0]
        result.append({**dict(p), 'num_fotos': cnt, 'num_revisiones': rev_cnt})
    conn.close()
    return jsonify(result)


@app.route('/api/pacientes', methods=['POST'])
def create_paciente():
    d = request.json
    conn = get_db()
    cur = conn.execute(
        '''INSERT INTO pacientes
           (nombre, fecha_nac, sexo, dni, telefono, email,
            objetivo, actividad_fisica, intolerancias, patologias, notas, peso, talla)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)''',
        (d['nombre'].strip(), d.get('fecha_nac') or None, d.get('sexo') or None,
         d.get('dni') or None, d.get('telefono') or None, d.get('email') or None,
         d.get('objetivo') or None, d.get('actividad_fisica') or None,
         d.get('intolerancias') or None, d.get('patologias') or None,
         d.get('notas') or None, d.get('peso') or None, d.get('talla') or None)
    )
    conn.commit()
    p = conn.execute('SELECT * FROM pacientes WHERE id=?', (cur.lastrowid,)).fetchone()
    conn.close()
    return jsonify({**dict(p), 'num_fotos': 0, 'num_revisiones': 0}), 201


@app.route('/api/pacientes/<int:pid>', methods=['PUT'])
def update_paciente(pid):
    d = request.json
    conn = get_db()
    conn.execute(
        '''UPDATE pacientes SET nombre=?, fecha_nac=?, sexo=?, dni=?, telefono=?, email=?,
           objetivo=?, actividad_fisica=?, intolerancias=?, patologias=?, notas=?, peso=?, talla=?
           WHERE id=?''',
        (d['nombre'].strip(), d.get('fecha_nac') or None, d.get('sexo') or None,
         d.get('dni') or None, d.get('telefono') or None, d.get('email') or None,
         d.get('objetivo') or None, d.get('actividad_fisica') or None,
         d.get('intolerancias') or None, d.get('patologias') or None,
         d.get('notas') or None, d.get('peso') or None, d.get('talla') or None, pid)
    )
    conn.commit()
    p       = conn.execute('SELECT * FROM pacientes WHERE id=?', (pid,)).fetchone()
    cnt     = conn.execute('SELECT COUNT(*) FROM fotos WHERE paciente_id=?', (pid,)).fetchone()[0]
    rev_cnt = conn.execute('SELECT COUNT(*) FROM revisiones WHERE paciente_id=?', (pid,)).fetchone()[0]
    conn.close()
    return jsonify({**dict(p), 'num_fotos': cnt, 'num_revisiones': rev_cnt})


@app.route('/api/pacientes/<int:pid>', methods=['DELETE'])
def delete_paciente(pid):
    conn = get_db()
    p = conn.execute('SELECT nombre FROM pacientes WHERE id=?', (pid,)).fetchone()
    if p:
        carpeta = os.path.join(FOTOS_DIR, nombre_a_carpeta(p['nombre']))
        if os.path.exists(carpeta):
            shutil.rmtree(carpeta)
    conn.execute('DELETE FROM fotos WHERE paciente_id=?', (pid,))
    conn.execute('DELETE FROM revisiones WHERE paciente_id=?', (pid,))
    conn.execute('DELETE FROM pacientes WHERE id=?', (pid,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


# ─── REVISIONES ───────────────────────────────────────────────────────────────

REV_FIELDS = [
    'fecha', 'peso', 'talla', 'envergadura',
    'pl_bicipital', 'pl_tricipital', 'pl_subescapular', 'pl_supraespinal',
    'pl_supracrestal', 'pl_abdominal', 'pl_muslo_int', 'pl_pierna_med',
    'per_antebrazo', 'per_brazo_rel', 'per_brazo_cont', 'per_hombros',
    'per_pecho', 'per_cintura', 'per_caderas', 'per_muslo', 'per_muslo_cont',
    'per_gemelo', 'notas'
]

def _rev_val(d, f):
    v = d.get(f)
    return v if v is not None and v != '' else None

def _rev_vals(d):
    return [_rev_val(d, f) for f in REV_FIELDS]


@app.route('/api/pacientes/<int:pid>/revisiones', methods=['GET'])
def get_revisiones(pid):
    conn = get_db()
    rows = conn.execute(
        'SELECT * FROM revisiones WHERE paciente_id=? ORDER BY fecha DESC', (pid,)
    ).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.route('/api/pacientes/<int:pid>/revisiones', methods=['POST'])
def create_revision(pid):
    d    = request.json
    cols = ', '.join(REV_FIELDS)
    ph   = ', '.join(['?'] * len(REV_FIELDS))
    conn = get_db()
    cur  = conn.execute(
        f'INSERT INTO revisiones (paciente_id, {cols}) VALUES (?, {ph})',
        [pid] + _rev_vals(d)
    )
    conn.commit()
    r = conn.execute('SELECT * FROM revisiones WHERE id=?', (cur.lastrowid,)).fetchone()
    conn.close()
    return jsonify(dict(r)), 201


@app.route('/api/revisiones/<int:rid>', methods=['PUT'])
def update_revision(rid):
    d    = request.json
    sets = ', '.join(f'{f}=?' for f in REV_FIELDS)
    conn = get_db()
    conn.execute(f'UPDATE revisiones SET {sets} WHERE id=?', _rev_vals(d) + [rid])
    conn.commit()
    r = conn.execute('SELECT * FROM revisiones WHERE id=?', (rid,)).fetchone()
    conn.close()
    return jsonify(dict(r))


@app.route('/api/revisiones/<int:rid>', methods=['DELETE'])
def delete_revision(rid):
    conn = get_db()
    conn.execute('DELETE FROM revisiones WHERE id=?', (rid,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


# ─── FOTOS ────────────────────────────────────────────────────────────────────

@app.route('/api/pacientes/<int:pid>/fotos', methods=['GET'])
def get_fotos(pid):
    conn = get_db()
    fotos = conn.execute(
        'SELECT * FROM fotos WHERE paciente_id=? ORDER BY fecha DESC', (pid,)
    ).fetchall()
    conn.close()
    return jsonify([dict(f) for f in fotos])


@app.route('/api/pacientes/<int:pid>/fotos', methods=['POST'])
def upload_foto(pid):
    file = request.files.get('foto')
    nota = request.form.get('nota', '').strip()
    if not file:
        return jsonify({'error': 'Sin archivo'}), 400
    conn     = get_db()
    p        = conn.execute('SELECT nombre FROM pacientes WHERE id=?', (pid,)).fetchone()
    carpeta  = nombre_a_carpeta(p['nombre'])
    prefijo  = nombre_a_prefijo(p['nombre'])
    ext      = os.path.splitext(file.filename)[1].lower() or '.jpg'
    fecha_str = datetime.now().strftime('%Y%m%d_%H%M%S')
    fname    = f"{prefijo}_{fecha_str}_{uuid.uuid4().hex[:6]}{ext}"
    filename = f"{carpeta}/{fname}"
    os.makedirs(os.path.join(FOTOS_DIR, carpeta), exist_ok=True)
    file.save(os.path.join(FOTOS_DIR, carpeta, fname))
    cur = conn.execute(
        'INSERT INTO fotos (paciente_id, filename, nota) VALUES (?,?,?)',
        (pid, filename, nota or None)
    )
    conn.commit()
    foto = conn.execute('SELECT * FROM fotos WHERE id=?', (cur.lastrowid,)).fetchone()
    conn.close()
    return jsonify(dict(foto)), 201


@app.route('/api/fotos/<int:fid>', methods=['DELETE'])
def delete_foto(fid):
    conn = get_db()
    f = conn.execute('SELECT * FROM fotos WHERE id=?', (fid,)).fetchone()
    if f:
        path = os.path.join(FOTOS_DIR, f['filename'].replace('/', os.sep))
        if os.path.exists(path):
            os.remove(path)
        conn.execute('DELETE FROM fotos WHERE id=?', (fid,))
        conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/fotos/<path:filename>')
def serve_foto(filename):
    return send_from_directory(FOTOS_DIR, filename)


# ─── EXPORTAR EXCEL ───────────────────────────────────────────────────────────

@app.route('/api/exportar-excel')
def exportar_excel():
    conn = get_db()
    pacientes = conn.execute('SELECT * FROM pacientes ORDER BY nombre COLLATE NOCASE').fetchall()

    wb = openpyxl.Workbook()

    verde  = PatternFill(start_color='16A34A', end_color='16A34A', fill_type='solid')
    blanco = Font(bold=True, color='FFFFFF', size=11)
    centro = Alignment(horizontal='center', vertical='center')

    def cabecera(ws, cabs, anchos):
        for i, (cab, ancho) in enumerate(zip(cabs, anchos), 1):
            cell = ws.cell(row=1, column=i, value=cab)
            cell.fill = verde; cell.font = blanco; cell.alignment = centro
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho
        ws.row_dimensions[1].height = 22

    # ── Hoja 1: Pacientes ──────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = 'Pacientes'
    cabecera(ws1,
        ['ID', 'Nombre', 'Fecha nac.', 'Sexo', 'DNI', 'Teléfono', 'Email',
         'Objetivo', 'Actividad física', 'Intolerancias', 'Patologías',
         'Notas', 'Nº revisiones', 'Nº fotos', 'Fecha alta'],
        [6, 28, 14, 10, 14, 14, 28, 22, 20, 30, 30, 35, 12, 8, 18]
    )
    for p in pacientes:
        cnt     = conn.execute('SELECT COUNT(*) FROM fotos WHERE paciente_id=?', (p['id'],)).fetchone()[0]
        rev_cnt = conn.execute('SELECT COUNT(*) FROM revisiones WHERE paciente_id=?', (p['id'],)).fetchone()[0]
        ws1.append([
            p['id'], p['nombre'], p['fecha_nac'] or '', p['sexo'] or '',
            p['dni'] or '', p['telefono'] or '', p['email'] or '',
            p['objetivo'] or '', p['actividad_fisica'] or '',
            p['intolerancias'] or '', p['patologias'] or '',
            p['notas'] or '', rev_cnt, cnt, p['creado'] or ''
        ])

    # ── Hoja 2: Revisiones ─────────────────────────────────────────────────
    ws2 = wb.create_sheet('Revisiones')
    cabecera(ws2,
        ['ID Rev.', 'Paciente', 'Fecha revisión',
         'Peso (kg)', 'Talla (cm)', 'Envergadura (cm)',
         'Pl. Bicipital', 'Pl. Tricipital', 'Pl. Subescapular', 'Pl. Supraespinal',
         'Pl. Supracrestal', 'Pl. Abdominal', 'Pl. Muslo int.', 'Pl. Pierna med.',
         'Per. Antebrazo', 'Per. Brazo rel.', 'Per. Brazo cont.', 'Per. Hombros',
         'Per. Pecho', 'Per. Cintura', 'Per. Caderas', 'Per. Muslo',
         'Per. Muslo cont.', 'Per. Gemelo', 'Notas'],
        [8, 28, 16, 10, 10, 14,
         13, 13, 16, 16, 16, 14, 14, 15,
         14, 14, 15, 13, 11, 12, 12, 11, 14, 12, 35]
    )
    revisiones = conn.execute(
        '''SELECT r.*, p.nombre FROM revisiones r
           JOIN pacientes p ON p.id = r.paciente_id
           ORDER BY p.nombre COLLATE NOCASE, r.fecha DESC'''
    ).fetchall()
    for r in revisiones:
        def v(f): return r[f] if r[f] is not None else ''
        ws2.append([
            v('id'), v('nombre'), v('fecha'),
            v('peso'), v('talla'), v('envergadura'),
            v('pl_bicipital'), v('pl_tricipital'), v('pl_subescapular'), v('pl_supraespinal'),
            v('pl_supracrestal'), v('pl_abdominal'), v('pl_muslo_int'), v('pl_pierna_med'),
            v('per_antebrazo'), v('per_brazo_rel'), v('per_brazo_cont'), v('per_hombros'),
            v('per_pecho'), v('per_cintura'), v('per_caderas'), v('per_muslo'),
            v('per_muslo_cont'), v('per_gemelo'), v('notas')
        ])

    conn.close()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nombre = f"nutricion_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True, download_name=nombre)


# ─── ARRANQUE ─────────────────────────────────────────────────────────────────

if __name__ == '__main__':
    init_db()
    print('\n  NutriClínica corriendo en http://localhost:5001\n')
    app.run(debug=False, port=5001)
